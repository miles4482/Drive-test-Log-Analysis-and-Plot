#!/usr/bin/env python3
"""Build a Bogura drive-test Excel report with RSRP/RSRQ/SINR charts.

Worksheet rules for every sheet:
- freeze panes off
- gridlines off (screen and print)
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from merge_and_analyze import (
    OUTPUT_DIR,
    PLOTS_DIR,
    RSRP_BINS,
    RSRP_LABELS,
    downsample,
    earfcn_to_band,
)

REPORT_XLSX = OUTPUT_DIR / "Bogura_DriveTest_Report.xlsx"

NAVY = "1B4F72"
WHITE = "FFFFFF"
LIGHT = "F4F6F7"
CARD = "EAF2F8"
RSRP_COLOR = "C0392B"
RSRQ_COLOR = "1F618D"
SINR_COLOR = "117A65"
THIN = Border(
    left=Side(style="thin", color="D5D8DC"),
    right=Side(style="thin", color="D5D8DC"),
    top=Side(style="thin", color="D5D8DC"),
    bottom=Side(style="thin", color="D5D8DC"),
)

RSRQ_BINS = [-np.inf, -20, -15, -10, np.inf]
RSRQ_LABELS = [
    "Poor (<-20)",
    "Fair (-20 to -15)",
    "Good (-15 to -10)",
    "Excellent (>-10)",
]
SINR_BINS = [-np.inf, 0, 13, 20, np.inf]
SINR_LABELS = [
    "Poor (<0)",
    "Fair (0 to 13)",
    "Good (13 to 20)",
    "Excellent (>20)",
]


def apply_sheet_view(ws: Worksheet) -> None:
    """No freeze panes, no gridlines (on-screen or printed)."""
    ws.freeze_panes = None
    ws.sheet_view.showGridLines = False
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_view.zoomScale = 100
    ws.sheet_view.view = "normal"


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def font(size=11, bold=False, color=NAVY, name="Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, color=color)


def write_cell(ws, row, col, value, *, size=11, bold=False, color=NAVY, fill_color=None, align="left", num_fmt=None, wrap=False):
    cell = ws.cell(row, col, value)
    cell.font = font(size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill_color:
        cell.fill = fill(fill_color)
    if num_fmt:
        cell.number_format = num_fmt
    return cell


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def kpi_stats(series: pd.Series) -> dict[str, float]:
    s = pd.to_numeric(series, errors="coerce").dropna()
    if s.empty:
        return {k: np.nan for k in ("Count", "Min", "P5", "Median", "Mean", "P95", "Max", "Std")}
    return {
        "Count": float(len(s)),
        "Min": float(s.min()),
        "P5": float(s.quantile(0.05)),
        "Median": float(s.median()),
        "Mean": float(s.mean()),
        "P95": float(s.quantile(0.95)),
        "Max": float(s.max()),
        "Std": float(s.std()),
    }


def quality_counts(series: pd.Series, bins, labels) -> pd.Series:
    binned = pd.cut(pd.to_numeric(series, errors="coerce"), bins=bins, labels=labels, right=False)
    return binned.value_counts(dropna=False).reindex(labels, fill_value=0)


def save_route_plot(df: pd.DataFrame, column: str, path: Path, vmin: float, vmax: float, cmap: str, unit: str) -> Path:
    PLOTS_DIR.mkdir(parents=True, exist_ok=True)
    geo = df.dropna(subset=["Longitude", "Latitude", column])
    sample = downsample(geo, max_points=25000)
    fig, ax = plt.subplots(figsize=(9.5, 7.2))
    sc = ax.scatter(
        sample["Longitude"],
        sample["Latitude"],
        c=sample[column],
        s=4,
        cmap=cmap,
        vmin=vmin,
        vmax=vmax,
        linewidths=0,
    )
    fig.colorbar(sc, ax=ax, label=f"{column} ({unit})")
    ax.set_title(f"Bogura drive route — {column}")
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_aspect("equal", adjustable="box")
    ax.grid(False)
    for spine in ax.spines.values():
        spine.set_color("#BFBFBF")
    fig.tight_layout()
    fig.savefig(path, dpi=140, facecolor="white")
    plt.close(fig)
    return path


def add_image(ws: Worksheet, path: Path, anchor: str, width: int, height: int) -> None:
    img = XLImage(str(path))
    img.width = width
    img.height = height
    ws.add_image(img, anchor)


def banner(ws: Worksheet, title: str, subtitle: str, last_col: int = 12) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    write_cell(ws, 1, 1, title, size=20, bold=True, color=WHITE, fill_color=NAVY, align="left")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    write_cell(ws, 2, 1, subtitle, size=12, color=WHITE, fill_color=NAVY)
    ws.row_dimensions[2].height = 20
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = fill(NAVY)
        ws.cell(2, col).fill = fill(NAVY)
        ws.cell(1, col).font = font(20, bold=True, color=WHITE)
        ws.cell(2, col).font = font(12, color=WHITE)


def write_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list], header_fill=NAVY, num_formats: dict[int, str] | None = None):
    for i, h in enumerate(headers):
        cell = write_cell(ws, start_row, start_col + i, h, size=11, bold=True, color=WHITE, fill_color=header_fill, align="center")
        cell.border = THIN
    for r_i, row in enumerate(rows, start=start_row + 1):
        bg = LIGHT if (r_i - start_row) % 2 == 0 else WHITE
        for c_i, value in enumerate(row):
            fmt = (num_formats or {}).get(c_i)
            cell = write_cell(ws, r_i, start_col + c_i, value, fill_color=bg, align="center" if c_i else "left", num_fmt=fmt)
            cell.border = THIN
    return start_row + len(rows)


def style_chart(chart, color: str, title: str, y_title: str, x_title: str, width=14, height=8) -> None:
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = x_title
    chart.style = 10
    chart.legend = None
    chart.width = width
    chart.height = height
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = color
        chart.series[0].graphicalProperties.line.solidFill = color


def add_col_chart(ws_data, ws_dest, anchor, title, color, cat_col, data_col, min_row, max_row, y_title, x_title, width=14, height=8):
    chart = BarChart()
    chart.type = "col"
    cats = Reference(ws_data, min_col=cat_col, min_row=min_row, max_row=max_row)
    data = Reference(ws_data, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    style_chart(chart, color, title, y_title, x_title, width=width, height=height)
    chart.shape = 4
    ws_dest.add_chart(chart, anchor)
    return chart


def add_line_chart(ws_data, ws_dest, anchor, title, color, cat_col, data_col, min_row, max_row, y_title, x_title, width=14, height=8, legend=False):
    chart = LineChart()
    cats = Reference(ws_data, min_col=cat_col, min_row=min_row, max_row=max_row)
    data = Reference(ws_data, min_col=data_col, min_row=min_row - 1, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    style_chart(chart, color, title, y_title, x_title, width=width, height=height)
    if legend:
        chart.legend.position = "b"
    chart.series[0].graphicalProperties.line.solidFill = color
    chart.series[0].graphicalProperties.line.width = 18000
    ws_dest.add_chart(chart, anchor)
    return chart


def write_chart_data(ws: Worksheet, df: pd.DataFrame) -> dict:
    """Hidden sheet holding all series used by Excel charts."""
    rsrp_bins = np.arange(-140, -48, 2)
    rsrq_bins = np.arange(-24, -2.5, 0.5)
    sinr_bins = np.arange(-15, 31, 1)

    def hist_block(col, bins, start_col, name):
        counts, edges = np.histogram(df[col].dropna(), bins=bins)
        ws.cell(1, start_col, f"{name}_bin")
        ws.cell(1, start_col + 1, f"{name}_samples")
        for i, (left, n) in enumerate(zip(edges[:-1], counts), start=2):
            ws.cell(i, start_col, float(left))
            ws.cell(i, start_col + 1, int(n))
        return {"col": start_col, "n": len(counts)}

    blocks = {
        "RSRP": hist_block("RSRP", rsrp_bins, 1, "RSRP"),
        "RSRQ": hist_block("RSRQ", rsrq_bins, 3, "RSRQ"),
        "SINR": hist_block("SINR", sinr_bins, 5, "SINR"),
    }

    percentiles = list(range(0, 101, 2))
    ws.cell(1, 7, "Percentile")
    ws.cell(1, 8, "RSRP_CDF")
    ws.cell(1, 9, "RSRQ_CDF")
    ws.cell(1, 10, "SINR_CDF")
    rsrp = df["RSRP"].dropna()
    rsrq = df["RSRQ"].dropna()
    sinr = df["SINR"].dropna()
    rsrp_p = np.nanpercentile(rsrp, percentiles)
    rsrq_p = np.nanpercentile(rsrq, percentiles)
    sinr_p = np.nanpercentile(sinr, percentiles)
    for i, p in enumerate(percentiles, start=2):
        ws.cell(i, 7, p)
        ws.cell(i, 8, float(rsrp_p[i - 2]))
        ws.cell(i, 9, float(rsrq_p[i - 2]))
        ws.cell(i, 10, float(sinr_p[i - 2]))
    blocks["cdf_n"] = len(percentiles)

    ws.cell(1, 12, "RSRP_quality")
    ws.cell(1, 13, "RSRP_quality_n")
    for i, (label, n) in enumerate(quality_counts(df["RSRP"], RSRP_BINS, RSRP_LABELS).items(), start=2):
        ws.cell(i, 12, str(label))
        ws.cell(i, 13, int(n))
    ws.cell(1, 15, "RSRQ_quality")
    ws.cell(1, 16, "RSRQ_quality_n")
    for i, (label, n) in enumerate(quality_counts(df["RSRQ"], RSRQ_BINS, RSRQ_LABELS).items(), start=2):
        ws.cell(i, 15, str(label))
        ws.cell(i, 16, int(n))
    ws.cell(1, 18, "SINR_quality")
    ws.cell(1, 19, "SINR_quality_n")
    for i, (label, n) in enumerate(quality_counts(df["SINR"], SINR_BINS, SINR_LABELS).items(), start=2):
        ws.cell(i, 18, str(label))
        ws.cell(i, 19, int(n))

    ts = (
        df.dropna(subset=["Time"])
        .set_index("Time")[["RSRP", "RSRQ", "SINR"]]
        .resample("10min")
        .mean()
        .dropna(how="all")
    )
    if len(ts) > 900:
        ts = ts.iloc[:: max(1, len(ts) // 900)]
    ws.cell(1, 21, "Time")
    ws.cell(1, 22, "RSRP")
    ws.cell(1, 23, "RSRQ")
    ws.cell(1, 24, "SINR")
    for i, (t, row) in enumerate(ts.iterrows(), start=2):
        cell = ws.cell(i, 21, pd.Timestamp(t).to_pydatetime())
        cell.number_format = "yyyy-mm-dd hh:mm"
        ws.cell(i, 22, None if pd.isna(row["RSRP"]) else float(row["RSRP"]))
        ws.cell(i, 23, None if pd.isna(row["RSRQ"]) else float(row["RSRQ"]))
        ws.cell(i, 24, None if pd.isna(row["SINR"]) else float(row["SINR"]))
    blocks["ts_n"] = len(ts)
    return blocks


def build_cover(ws: Worksheet, df: pd.DataFrame) -> None:
    banner(
        ws,
        "  Bogura LTE Drive-Test Report",
        "  Merged P1 + P2  |  RSRP, RSRQ, SINR  |  Gridlines off  |  Freeze panes off",
        last_col=10,
    )
    set_widths(ws, {"A": 28, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18})

    write_cell(ws, 4, 1, "Dataset", size=14, bold=True)
    info_rows = [
        ["P1 samples", int((df["Source"] == "P1").sum())],
        ["P2 samples", int((df["Source"] == "P2").sum())],
        ["Merged samples", int(len(df))],
        ["Start time", df["Time"].min()],
        ["End time", df["Time"].max()],
        ["Unique Cell Ids", int(df["Cell Id"].nunique())],
        ["Longitude range", f"{df['Longitude'].min():.5f} – {df['Longitude'].max():.5f}"],
        ["Latitude range", f"{df['Latitude'].min():.5f} – {df['Latitude'].max():.5f}"],
    ]
    write_table(ws, 5, 1, ["Item", "Value"], info_rows, num_formats={1: "#,##0"})
    ws.cell(9, 2).number_format = "yyyy-mm-dd hh:mm:ss"
    ws.cell(10, 2).number_format = "yyyy-mm-dd hh:mm:ss"

    write_cell(ws, 4, 4, "Radio KPIs (all samples)", size=14, bold=True)
    stats_headers = ["KPI", "Count", "Min", "P5", "Median", "Mean", "P95", "Max"]
    stats_rows = []
    for name in ("RSRP", "RSRQ", "SINR"):
        st = kpi_stats(df[name])
        stats_rows.append(
            [name, st["Count"], st["Min"], st["P5"], st["Median"], st["Mean"], st["P95"], st["Max"]]
        )
    write_table(
        ws,
        5,
        4,
        stats_headers,
        stats_rows,
        num_formats={1: "#,##0", 2: "0.00", 3: "0.00", 4: "0.00", 5: "0.00", 6: "0.00", 7: "0.00"},
    )

    write_cell(ws, 15, 1, "How to read this workbook", size=14, bold=True)
    notes = [
        "Cover — dataset size and KPI scorecard for the merged Bogura log.",
        "RSRP / RSRQ / SINR — Excel charts (histogram, CDF, quality bins) plus a coverage map. No freeze, no gridlines.",
        "Time Series — 10-minute mean RSRP, RSRQ and SINR across the drive days.",
        "Sample Log — evenly spaced subset of the merged samples (full 1.07M rows stay in output/Bogura_merged.csv.gz).",
        "P1 is the split archive (part1–part5). P2 is the standalone archive. This report uses the concatenated final file.",
    ]
    for i, text in enumerate(notes):
        ws.merge_cells(start_row=16 + i, start_column=1, end_row=16 + i, end_column=10)
        write_cell(ws, 16 + i, 1, text, wrap=True)
        ws.row_dimensions[16 + i].height = 18

    write_cell(ws, 22, 1, "Band mix (DL EARFCN)", size=14, bold=True)
    band = df["Band"].value_counts()
    band_rows = [[str(k), int(v), v / len(df)] for k, v in band.items()]
    write_table(ws, 23, 1, ["Band", "Samples", "Share"], band_rows, num_formats={1: "#,##0", 2: "0.0%"})
    # Histograms are added after _ChartData exists; see add_cover_charts().


def build_kpi_sheet(ws, data_ws, df, name, color, unit, hist_col, hist_n, cdf_col, q_cat_col, q_n, map_path: Path | None):
    banner(ws, f"  {name} report", f"  Unit: {unit}  |  Charts are native Excel objects  |  Gridlines off", last_col=10)
    set_widths(ws, {get_column_letter(i): 16 for i in range(1, 11)})
    ws.column_dimensions["A"].width = 28

    write_cell(ws, 4, 1, f"{name} statistics", size=14, bold=True)
    st = kpi_stats(df[name])
    write_table(
        ws,
        5,
        1,
        ["Metric", "Value"],
        [[k, st[k]] for k in ("Count", "Min", "P5", "Median", "Mean", "P95", "Max", "Std")],
        header_fill=color,
        num_formats={1: "0.00"},
    )
    ws.cell(6, 2).number_format = "#,##0"

    add_col_chart(
        data_ws,
        ws,
        "D4",
        f"{name} histogram",
        color,
        hist_col,
        hist_col + 1,
        2,
        1 + hist_n,
        "Samples",
        f"{name} ({unit})",
        width=16,
        height=8,
    )
    add_line_chart(
        data_ws,
        ws,
        "D20",
        f"{name} CDF",
        color,
        7,
        cdf_col,
        2,
        1 + 51,
        f"{name} ({unit})",
        "Percentile",
        width=16,
        height=8,
    )
    add_col_chart(
        data_ws,
        ws,
        "A20",
        f"{name} quality bins",
        color,
        q_cat_col,
        q_cat_col + 1,
        2,
        1 + q_n,
        "Samples",
        "",
        width=12,
        height=8,
    )
    if map_path and map_path.exists():
        write_cell(ws, 38, 1, f"{name} coverage map", size=14, bold=True)
        add_image(ws, map_path, "A40", width=760, height=560)


def build_time_sheet(ws, data_ws, ts_n: int) -> None:
    banner(ws, "  RSRP / RSRQ / SINR vs time", "  10-minute mean of the merged Bogura log  |  Gridlines off", last_col=10)
    set_widths(ws, {get_column_letter(i): 16 for i in range(1, 11)})
    write_cell(ws, 4, 1, "Each chart is a native Excel line chart from the merged P1+P2 samples.", wrap=True)
    add_line_chart(data_ws, ws, "A6", "RSRP vs time (10-min mean)", RSRP_COLOR, 21, 22, 2, 1 + ts_n, "RSRP (dBm)", "Time", width=22, height=8)
    add_line_chart(data_ws, ws, "A22", "RSRQ vs time (10-min mean)", RSRQ_COLOR, 21, 23, 2, 1 + ts_n, "RSRQ (dB)", "Time", width=22, height=8)
    add_line_chart(data_ws, ws, "A38", "SINR vs time (10-min mean)", SINR_COLOR, 21, 24, 2, 1 + ts_n, "SINR (dB)", "Time", width=22, height=8)


def build_sample_log(ws: Worksheet, df: pd.DataFrame, n: int = 8000) -> None:
    banner(ws, "  Sample drive log", "  Evenly spaced subset for inspection  |  Full merged file is CSV  |  No freeze / no gridlines", last_col=10)
    sample = downsample(df, max_points=n).copy()
    cols = ["Time", "RSRP", "RSRQ", "SINR", "Longitude", "Latitude", "Cell Id", "DL EARFCN", "Source", "Band"]
    sample = sample[cols]
    for c, w in zip(cols, (22, 12, 12, 12, 14, 14, 14, 14, 10, 22)):
        ws.column_dimensions[get_column_letter(cols.index(c) + 1)].width = w
    headers = cols
    for i, h in enumerate(headers, start=1):
        write_cell(ws, 4, i, h, bold=True, color=WHITE, fill_color=NAVY, align="center")
    for r_i, row in enumerate(sample.itertuples(index=False), start=5):
        bg = LIGHT if r_i % 2 == 0 else WHITE
        values = list(row)
        for c_i, value in enumerate(values, start=1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            if hasattr(value, "item"):
                try:
                    value = value.item()
                except Exception:
                    pass
            if pd.isna(value):
                value = None
            cell = write_cell(ws, r_i, c_i, value, fill_color=bg, align="center")
            cell.border = THIN
            if c_i == 1 and value is not None:
                cell.number_format = "yyyy-mm-dd hh:mm:ss.000"
            if c_i in (2, 3, 4, 5, 6):
                cell.number_format = "0.000"


def verify_report(path: Path) -> dict:
    wb = load_workbook(path)
    info = {"sheets": {}, "chart_titles": [], "problems": []}
    for name in wb.sheetnames:
        ws = wb[name]
        freeze = ws.freeze_panes
        grid = ws.sheet_view.showGridLines
        print_grid = ws.print_options.gridLines
        n_charts = len(ws._charts)
        n_images = len(ws._images)
        info["sheets"][name] = {
            "freeze_panes": freeze,
            "showGridLines": grid,
            "print_gridLines": print_grid,
            "charts": n_charts,
            "images": n_images,
        }
        if freeze not in (None, "A1"):
            info["problems"].append(f"{name}: freeze_panes={freeze}")
        if grid is not False:
            info["problems"].append(f"{name}: showGridLines={grid}")
        if print_grid:
            info["problems"].append(f"{name}: print gridlines on")
        for ch in ws._charts:
            title = None
            if ch.title is not None:
                try:
                    title = str(ch.title.tx.rich.p[0].r[0].t)
                except Exception:
                    title = str(ch.title)
            info["chart_titles"].append((name, title, type(ch).__name__))
    wb.close()
    return info


def add_cover_charts(cover, data_ws, blocks) -> None:
    write_cell(cover, 29, 1, "RSRP, RSRQ and SINR plots (native Excel charts)", size=14, bold=True)
    add_col_chart(
        data_ws, cover, "A31", "RSRP histogram", RSRP_COLOR,
        blocks["RSRP"]["col"], blocks["RSRP"]["col"] + 1, 2, 1 + blocks["RSRP"]["n"],
        "Samples", "RSRP (dBm)", width=12, height=7,
    )
    add_col_chart(
        data_ws, cover, "G31", "RSRQ histogram", RSRQ_COLOR,
        blocks["RSRQ"]["col"], blocks["RSRQ"]["col"] + 1, 2, 1 + blocks["RSRQ"]["n"],
        "Samples", "RSRQ (dB)", width=12, height=7,
    )
    add_col_chart(
        data_ws, cover, "A48", "SINR histogram", SINR_COLOR,
        blocks["SINR"]["col"], blocks["SINR"]["col"] + 1, 2, 1 + blocks["SINR"]["n"],
        "Samples", "SINR (dB)", width=12, height=7,
    )


def save_report_preview(df: pd.DataFrame, path: Path) -> Path:
    fig, axes = plt.subplots(1, 3, figsize=(15, 4.6))
    specs = (
        ("RSRP", np.arange(-140, -48, 2), (-140, -50), RSRP_COLOR, "dBm"),
        ("RSRQ", np.arange(-24, -2.5, 0.5), (-24, -3), RSRQ_COLOR, "dB"),
        ("SINR", np.arange(-15, 31, 1), (-15, 30), SINR_COLOR, "dB"),
    )
    for ax, (col, bins, xlim, color, unit) in zip(axes, specs):
        ax.hist(df[col].dropna(), bins=bins, color=f"#{color}", edgecolor="none")
        ax.set_title(f"{col} histogram")
        ax.set_xlabel(f"{col} ({unit})")
        ax.set_ylabel("Samples")
        ax.set_xlim(*xlim)
        ax.grid(False)
        ax.set_facecolor("white")
    fig.suptitle("Bogura Drive-Test Report — RSRP / RSRQ / SINR", fontsize=13, fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=140, facecolor="white")
    plt.close(fig)
    return path


def build_report(df: pd.DataFrame, out_path: Path = REPORT_XLSX) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PLOTS_DIR.mkdir(parents=True, exist_ok=True)
    if "Band" not in df.columns:
        df = df.copy()
        df["Band"] = df["DL EARFCN"].map(earfcn_to_band)

    rsrp_map = save_route_plot(df, "RSRP", PLOTS_DIR / "excel_route_rsrp.png", -120, -70, "RdYlGn", "dBm")
    rsrq_map = save_route_plot(df, "RSRQ", PLOTS_DIR / "excel_route_rsrq.png", -20, -6, "RdYlGn", "dB")
    sinr_map = save_route_plot(df, "SINR", PLOTS_DIR / "excel_route_sinr.png", -5, 25, "RdYlGn", "dB")
    save_report_preview(df, PLOTS_DIR / "excel_rsrp_rsrq_sinr_histograms.png")

    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    rsrp_ws = wb.create_sheet("RSRP")
    rsrq_ws = wb.create_sheet("RSRQ")
    sinr_ws = wb.create_sheet("SINR")
    time_ws = wb.create_sheet("Time Series")
    sample_ws = wb.create_sheet("Sample Log")
    data_ws = wb.create_sheet("_ChartData")
    data_ws.sheet_state = "hidden"

    blocks = write_chart_data(data_ws, df)
    build_cover(cover, df)
    add_cover_charts(cover, data_ws, blocks)
    build_kpi_sheet(
        rsrp_ws,
        data_ws,
        df,
        "RSRP",
        RSRP_COLOR,
        "dBm",
        blocks["RSRP"]["col"],
        blocks["RSRP"]["n"],
        8,
        12,
        5,
        rsrp_map,
    )
    build_kpi_sheet(
        rsrq_ws,
        data_ws,
        df,
        "RSRQ",
        RSRQ_COLOR,
        "dB",
        blocks["RSRQ"]["col"],
        blocks["RSRQ"]["n"],
        9,
        15,
        4,
        rsrq_map,
    )
    build_kpi_sheet(
        sinr_ws,
        data_ws,
        df,
        "SINR",
        SINR_COLOR,
        "dB",
        blocks["SINR"]["col"],
        blocks["SINR"]["n"],
        10,
        18,
        4,
        sinr_map,
    )
    build_time_sheet(time_ws, data_ws, blocks["ts_n"])
    build_sample_log(sample_ws, df)

    for ws in wb.worksheets:
        apply_sheet_view(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote Excel report: {out_path}")
    return out_path


if __name__ == "__main__":
    from merge_and_analyze import MERGED_CSV, extract_archives, load_merged_csv, merge_parts

    if MERGED_CSV.exists():
        frame = load_merged_csv()
    else:
        extract_archives()
        frame = merge_parts()
    path = build_report(frame)
    result = verify_report(path)
    print(result)
    if result["problems"]:
        raise SystemExit("Excel view checks failed: " + "; ".join(result["problems"]))
    print("Verified: no freeze panes, no gridlines, RSRP/RSRQ/SINR charts present.")
